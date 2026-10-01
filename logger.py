from openpyxl import Workbook, load_workbook
from datetime import datetime
from pathlib import Path
import os


if os.getenv("VERCEL"):
    CARPETA_MEDIFY = Path("/tmp/medify_logs")
else:
    DESCARGAS = Path.home() / "Downloads"
    CARPETA_MEDIFY = DESCARGAS / "MEDIFY_AI"


CARPETA_MEDIFY.mkdir(
    parents=True,
    exist_ok=True
)

ARCHIVO_LOG = CARPETA_MEDIFY / "LOG.xlsx"


def registrar_proceso(datos):

    CARPETA_MEDIFY.mkdir(
        parents=True,
        exist_ok=True
    )

    if ARCHIVO_LOG.exists():

        wb = load_workbook(ARCHIVO_LOG)
        ws = wb.active

    else:

        wb = Workbook()
        ws = wb.active

        ws.append([
            "Fecha",
            "Archivo original",
            "Identificacion",
            "Clasificacion",
            "Archivo generado",
            "Estado",
            "Detalle"
        ])

    ws.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        datos.get("archivo_original", ""),
        datos.get("identificacion", ""),
        datos.get("clasificacion", ""),
        datos.get("archivo_generado", ""),
        datos.get("estado", "OK"),
        datos.get("detalle", "")
    ])

    wb.save(ARCHIVO_LOG)

    print(
        f"LOG actualizado: {ARCHIVO_LOG}"
    )